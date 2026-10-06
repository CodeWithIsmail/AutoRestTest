import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x1F, 0x39, 0x64) # Navy blue
        if level == 2:
            run.font.size = Pt(14)
            run.bold = True
        elif level == 3:
            run.font.size = Pt(13)
            run.bold = True
        elif level == 4:
            run.font.size = Pt(12)
            run.bold = True
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.space_before = Pt(12)

def add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p

def set_table_header(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        shading_elm = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading_elm)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.bold = True

def add_table_row(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

def main():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    add_heading(doc, '4. Testing Components', 2)
    add_paragraph(doc, "This section outlines the Software Quality Assurance (SQA) procedures implemented to verify the stability, security, and functional correctness of the AutoRestTest platform itself (not the target APIs being tested).")
    
    add_heading(doc, '4.1 Testing Approach & Strategy', 3)
    add_paragraph(doc, "AutoRestTest employs a hybrid testing strategy designed to handle both deterministic web application logic and non-deterministic AI pipelines:")
    
    add_paragraph(doc, "1. Backend Unit Testing: ", bold=True)
    add_paragraph(doc, "The NestJS backend relies heavily on the Jest testing framework. Isolated unit tests (e.g., auth.service.spec.ts, endpoint-extractor.spec.ts) are executed by mocking the Prisma ORM layer, ensuring business logic and RBAC validations function correctly without hitting a live database.")
    
    add_paragraph(doc, "2. Frontend Validation: ", bold=True)
    add_paragraph(doc, "React components and TanStack Query state transitions are manually validated against edge cases, ensuring loading spinners, error toasts, and UI constraints (like missing permissions) trigger appropriately.")
    
    add_paragraph(doc, "3. AI Pipeline & Engine Testing: ", bold=True)
    add_paragraph(doc, "The Python MARL testing engine and LLM generation services utilize a \"Mock Engine Mode\" (ENGINE_MODE=mock). This bypasses live NVIDIA API calls, injecting static responses to deterministically verify the pipeline's handling of LLM outputs without incurring API costs or waiting for generation.")

    add_heading(doc, '4.2 Item Pass/Fail Criteria', 3)
    add_paragraph(doc, "To determine the success of an implemented feature, the following criteria are strictly applied:")
    add_paragraph(doc, "• Pass Criteria: The module executes without throwing unhandled exceptions. Expected database states (e.g., test runs recorded, specs parsed) exactly match the actual database state. Unit tests pass with a 0 exit code. For AI features, the mock engine completes the execution loop successfully.")
    add_paragraph(doc, "• Fail Criteria: The application crashes or hangs indefinitely (e.g., polling fails to terminate). Data corruption occurs. Unauthorized roles successfully access restricted endpoints. The UI fails to reflect the underlying state (e.g., test finishes but graph remains spinning).")

    add_heading(doc, '4.3 Risks and Contingencies', 3)
    
    table_risks = doc.add_table(rows=1, cols=3)
    table_risks.style = 'Table Grid'
    table_risks.autofit = False
    for i, w in enumerate([Inches(1.5), Inches(2.2), Inches(2.8)]):
        table_risks.columns[i].width = w
        for cell in table_risks.columns[i].cells: cell.width = w
    set_table_header(table_risks.rows[0], ['Risk Identified', 'Impact', 'Contingency Plan'])
    
    data_risks = [
        ("LLM Non-Determinism", "Automated tests relying on live LLM outputs may randomly fail due to varied responses.", "Utilize ENGINE_MODE=mock during testing to inject static, predictable JSON responses in place of live LLM generation."),
        ("Large Graph Rendering", "Cytoscape canvas may freeze the browser when rendering a massive enterprise API.", "Implement strict node-limits and utilize the Focus View to only render a 3-column subset of active nodes at any given time."),
        ("Polling Timeouts", "AI generation takes too long, causing the frontend to time out and display an error.", "Implement background job tracking. The frontend polls for status without holding an open connection, timing out gracefully if the engine halts.")
    ]
    for r in data_risks:
        add_table_row(table_risks.add_row(), r)
        
    add_paragraph(doc, "")

    add_heading(doc, '4.4 Test Cases with Detailed Outcomes', 3)
    add_paragraph(doc, "The table below details core integration and manual test cases executed against the platform.")

    table_tc = doc.add_table(rows=1, cols=4)
    table_tc.style = 'Table Grid'
    table_tc.autofit = False
    for i, w in enumerate([Inches(0.8), Inches(2.0), Inches(2.0), Inches(1.7)]):
        table_tc.columns[i].width = w
        for cell in table_tc.columns[i].cells: cell.width = w
    set_table_header(table_tc.rows[0], ['Test ID', 'Test Description (Scenario)', 'Expected Outcome', 'Actual Outcome'])

    data_tc = [
        ("TC-01", "User Registration OTP Verification: Submit valid registration, retrieve OTP, submit verification.", "User is successfully created, token stored, redirected to Dashboard.", "Pass"),
        ("TC-02", "Spec Parsing: Upload standard OpenAPI 3.0 YAML file to a project.", "Parser successfully extracts all endpoints into database and mounts them in the UI.", "Pass"),
        ("TC-03", "RBAC Tester Restriction: User with 'Tester' role attempts to delete an endpoint.", "Backend returns 403 Forbidden. Frontend hides delete button.", "Pass"),
        ("TC-04", "Graph Build Execution: Trigger a dependency graph build on 100+ endpoints.", "Engine resolves embeddings, UI polls until completion, canvas renders successfully.", "Pass"),
        ("TC-05", "Live Test Execution: Trigger MARL test suite generation on mock target.", "Status updates to 'Running', timeline streams passes/fails, final report is generated.", "Pass"),
        ("TC-06", "LLM Failure Explanation: Click 'AI Explain' on a failed HTTP request log.", "System prompts LLM with request/response context; plain text explanation streams to UI.", "Pass")
    ]
    for r in data_tc:
        add_table_row(table_tc.add_row(), r)

    output_path_root = os.path.join("D:\\SPL3\\AutoRestTest", "AutoRestTest_TestingComponents.docx")
    output_path_srs = os.path.join("D:\\SPL3\\AutoRestTest", "resources", "SRS", "AutoRestTest_TestingComponents.docx")
    
    doc.save(output_path_root)
    os.makedirs(os.path.dirname(output_path_srs), exist_ok=True)
    doc.save(output_path_srs)

if __name__ == "__main__":
    main()
