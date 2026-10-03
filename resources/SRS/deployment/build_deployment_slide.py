import copy, os
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRS = r'D:\SPL3\AutoRestTest\resources\SRS'
ICON = os.path.join(SRS, 'deployment', 'icons')
SRC = os.path.join(SRS, '1433_SPL3_FINAL.pptx')
OUT = os.path.join(SRS, 'deployment_slide.pptx')

BLUE = RGBColor(0x2F, 0x6F, 0xC6); DBLUE = RGBColor(0x1F, 0x49, 0x7D)
ORANGE = RGBColor(0xD9, 0x6B, 0x13); LIGHT = RGBColor(0xE4, 0xEC, 0xF7)
GREY = RGBColor(0x53, 0x5D, 0x6B); WHITE = RGBColor(255, 255, 255)
FONT = 'Merriweather Sans'

prs = Presentation(SRC)
# keep only slide 2 (Project Overview) as the style base
ids = prs.slides._sldIdLst
keep = list(ids)[1]
for sid in list(ids):
    if sid is not keep:
        prs.part.drop_rel(sid.get(qn('r:id')))
        ids.remove(sid)
slide = prs.slides[0]
shapes = list(slide.shapes)
for sh in shapes[4:]:           # keep bg + decorations + title only
    sh._element.getparent().remove(sh._element)
title = shapes[3]
title.text_frame.paragraphs[0].runs[0].text = 'Deployment'
for r in title.text_frame.paragraphs[0].runs[1:]:
    r.text = ''


def text(x, y, w, h, s, size, color, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = s
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb


def card(x, y, w, h, role, name, logos, accent):
    c = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    c.adjustments[0] = 0.08
    c.fill.solid(); c.fill.fore_color.rgb = WHITE
    c.line.color.rgb = LIGHT; c.line.width = Pt(2)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x + 330000), Emu(y), Emu(w - 660000), Emu(70000))
    bar.fill.solid(); bar.fill.fore_color.rgb = accent; bar.line.fill.background()
    text(x, y + 160000, w, 300000, role.upper(), 15, accent, True)
    n = len(logos); sz = 900000
    for i, (lg, nm) in enumerate(logos):
        slot = w / n
        cx = x + slot * i + slot / 2
        slide.shapes.add_picture(os.path.join(ICON, lg + '.png'), Emu(int(cx - sz / 2)), Emu(y + 560000), Emu(sz), Emu(sz))
        text(int(cx - slot / 2), y + 1560000, int(slot), 600000, nm, 24 if n == 1 else 20, DBLUE, True)
    return c


def arrow(x1, y1, x2, y2, label=None, lx=0, ly=0, lw=1500000):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Emu(x1), Emu(y1), Emu(x2), Emu(y2))
    cn.line.color.rgb = BLUE; cn.line.width = Pt(3.5)
    ln = cn.line._get_or_add_ln()
    ln.append(etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='lg', len='lg'))
    if label:
        text(lx, ly, lw, 320000, label, 14, GREY, True)


W, H, GAP = 4400000, 2400000, 1500000
X0 = (18288000 - (3 * W + 2 * GAP)) // 2
Y1, Y2 = 2700000, 6150000
cx = [X0 + i * (W + GAP) for i in range(3)]

card(cx[0], Y1, W, H, 'Frontend', 'Vercel', [('vercel', 'Vercel')], BLUE)
card(cx[1], Y1, W, H, 'Backend Service', 'Render', [('render', 'Render')], ORANGE)
card(cx[2], Y1, W, H, 'Database', 'Neon DB', [('neon', 'Neon DB')], BLUE)
card(cx[1], Y2, W, H, 'Engine Service', 'Render', [('render', 'Render')], ORANGE)
card(cx[2], Y2, W, H, 'LLM Providers', '', [('nvidia', 'NVIDIA NIM'), ('googlegemini', 'Google Gemini')], DBLUE)

mid1 = Y1 + H // 2
arrow(cx[0] + W, mid1, cx[1], mid1, 'REST API', cx[0] + W, mid1 - 420000)
arrow(cx[1] + W, mid1, cx[2], mid1, 'SQL', cx[1] + W, mid1 - 420000)
arrow(cx[1] + W // 2, Y1 + H, cx[1] + W // 2, Y2, 'REST API', cx[1] + W // 2 + 150000, Y1 + H + 330000, 1800000)
mid2 = Y2 + H // 2
arrow(cx[1] + W, mid2, cx[2], mid2, 'LLM API', cx[1] + W, mid2 - 420000)

# public URL card (same footprint as the other cards, points up at the frontend)
GLOBE = os.path.join(ICON, 'globe.png')
if not os.path.exists(GLOBE):
    import fitz
    svg = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 64 64" fill="none" stroke="#FFFFFF" '
           'stroke-width="3.5" stroke-linecap="round"><circle cx="32" cy="32" r="26"/>'
           '<ellipse cx="32" cy="32" rx="11" ry="26"/><path d="M6 32h52M10 18h44M10 46h44"/></svg>')
    pg = fitz.open(stream=svg.encode(), filetype='svg')[0]
    pg.get_pixmap(matrix=fitz.Matrix(8, 8), alpha=True).save(GLOBE)

cy = Y2
ban = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(cx[0]), Emu(cy), Emu(W), Emu(H))
ban.adjustments[0] = 0.08
ban.fill.gradient(); ban.fill.gradient_angle = 35
st = ban.fill.gradient_stops
st[0].color.rgb = DBLUE; st[0].position = 0
st[1].color.rgb = BLUE; st[1].position = 1
ban.line.fill.background()
ban.click_action.hyperlink.address = 'https://autoresttest.vercel.app'

# globe icon in a translucent circle
ci = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx[0] + 300000), Emu(cy + 280000), Emu(900000), Emu(900000))
ci.fill.solid(); ci.fill.fore_color.rgb = WHITE; ci.line.fill.background()
sf = ci.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
etree.SubElement(sf, qn('a:alpha'), val='22000')
slide.shapes.add_picture(GLOBE, Emu(cx[0] + 440000), Emu(cy + 420000), Emu(620000), Emu(620000))

# LIVE badge
bd = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(cx[0] + W - 1500000), Emu(cy + 430000), Emu(1200000), Emu(430000))
bd.adjustments[0] = 0.5
bd.fill.solid(); bd.fill.fore_color.rgb = WHITE; bd.line.fill.background()
dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Emu(cx[0] + W - 1380000), Emu(cy + 560000), Emu(170000), Emu(170000))
dot.fill.solid(); dot.fill.fore_color.rgb = RGBColor(0x1F, 0xB8, 0x4F); dot.line.fill.background()
text(cx[0] + W - 1150000, cy + 430000, 800000, 430000, 'LIVE', 15, DBLUE, True)

text(cx[0] + 1350000, cy + 400000, 1900000, 560000, 'PUBLIC URL', 17, LIGHT, True, PP_ALIGN.LEFT)

# URL pill
pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(cx[0] + 300000), Emu(cy + 1380000), Emu(W - 600000), Emu(700000))
pill.adjustments[0] = 0.5
pill.fill.solid(); pill.fill.fore_color.rgb = WHITE; pill.line.fill.background()
pill.click_action.hyperlink.address = 'https://autoresttest.vercel.app'
text(cx[0] + 300000, cy + 1380000, W - 600000, 700000, 'autoresttest.vercel.app', 22, BLUE, True)

arrow(cx[0] + W // 2, cy, cx[0] + W // 2, Y1 + H)

prs.save(OUT)
print('saved', OUT)
