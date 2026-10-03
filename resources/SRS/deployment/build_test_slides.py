import os, copy
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

SRS = r'D:\SPL3\AutoRestTest\resources\SRS'
SRC = os.path.join(SRS, '1433_SPL3_FINAL.pptx')
OUT = os.path.join(SRS, 'test_results_slides.pptx')

BLUE = RGBColor(0x2F, 0x6F, 0xC6); DBLUE = RGBColor(0x1F, 0x49, 0x7D)
ORANGE = RGBColor(0xD9, 0x6B, 0x13); LIGHT = RGBColor(0xE4, 0xEC, 0xF7)
GREY = RGBColor(0x53, 0x5D, 0x6B); WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(0x1F, 0xA8, 0x57); PURPLE = RGBColor(0x7A, 0x5C, 0xC8)
ROWALT = RGBColor(0xF4, 0xF7, 0xFC); GREENBG = RGBColor(0xE6, 0xF6, 0xEC)
FONT = 'Merriweather Sans'
I = 914400
def E(v): return Emu(int(v * I))

MODULES = [
    ('Authentication', 13, 'Email-confirmed signup, hashed credentials, password reset, account deletion'),
    ('Project Management', 3, 'Project creation, access control, safe deletion'),
    ('API Specification Mgmt', 12, 'Spec upload and validation, AI spec generation, endpoints, dependency graph'),
    ('Test Suite Management', 4, 'Run configuration, custom auth headers, endpoint exclusion'),
    ('Test Execution', 6, 'One-click async run, duplicate-run guard, re-run, request inspection'),
    ('Results & Reports', 5, 'Run reports, AI failure explanation, replay, history and comparison'),
    ('Team Collaboration', 7, 'Invitations with expiry, viewer / tester / admin role enforcement'),
]
MCOL = [BLUE, ORANGE, GREEN, PURPLE, BLUE, ORANGE, GREEN]

TITLES = """User Registration with Valid Details|User Registration with Invalid or Missing Fields|Registration with an Already Registered Email or Username|Re-registering an Unconfirmed Email Address|Email Confirmation with the Correct Code|Email Confirmation with a Wrong Code|Login with Wrong Password or Unknown Account|Password Reset|Password Reset for an Unknown Email|Update Display Name and Avatar Colour|Change Password|Notification Preferences|Delete Account|Create a Project|Project Access Control|Delete a Project|Upload a Valid OpenAPI 3.0 Specification|Upload a Malformed Specification|Replace an Existing Specification|Specification Permissions|Generate a Specification from a Source Archive|Reject an Invalid or Unsafe Source Upload|Review and Apply a Generated Specification|Add an Endpoint Manually|Delete an Endpoint|Build the Project Dependency Graph|Inspect One Operation's Dependencies|Dependency Graph Attached to a Completed Run|Configure a Test Run|Custom Authentication Headers|Custom Header Validation|Excluding Endpoints from a Run|One-Click Test Execution|Running Without a Specification, or While Already Running|Re-running an Existing Run|Per-Endpoint Request Summary|Browse and Filter Captured Requests|Inspect a Single Request and Response|Explain Requests|View a Run Report|Explain Failures|Replay a Completed Run|Run History and Comparison|Invite a Collaborator|Accept an Invitation|Invitation Security and Expiry|Resend, Revoke and Decline an Invitation|Role Enforcement: Viewer|Role Enforcement: Tester and Administrator|Manage Members""".split('|')

prs = Presentation(SRC)
ids = prs.slides._sldIdLst
keep = list(ids)[1]
for sid in list(ids):
    if sid is not keep:
        prs.part.drop_rel(sid.get(qn('r:id'))); ids.remove(sid)
tmpl = prs.slides[0]
tshapes = list(tmpl.shapes)
base = [copy.deepcopy(s._element) for s in tshapes[:4]]

def new_slide(title):
    s = prs.slides.add_slide(tmpl.slide_layout)
    for ph in list(s.shapes):
        ph._element.getparent().remove(ph._element)
    for el in base:
        e2 = copy.deepcopy(el)
        for bl in e2.iter(qn('a:blip')):
            rid = bl.get(qn('r:embed'))
            if rid:
                part = tmpl.part.related_part(rid)
                bl.set(qn('r:embed'), s.part.relate_to(part, RT.IMAGE))
        s.shapes._spTree.append(e2)
    tsh = list(s.shapes)[3]
    tsh.left, tsh.top, tsh.width = E(0.65), tsh.top, E(17.0)
    tsh.text_frame.word_wrap = True
    tsh.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    t = tsh.text_frame.paragraphs[0]
    t.runs[0].text = title
    for r in t.runs[1:]:
        r.text = ''
    return s

def text(s, x, y, w, h, txt, size, color, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tb = s.shapes.add_textbox(E(x), E(y), E(w), E(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = txt
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb

def box(s, x, y, w, h, fill=WHITE, line=LIGHT, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    b = s.shapes.add_shape(shape, E(x), E(y), E(w), E(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = radius
    b.fill.solid(); b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb = line; b.line.width = Pt(2)
    return b

def accent(s, x, y, w, col):
    a = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(x + 0.25), E(y), E(w - 0.5), E(0.06))
    a.fill.solid(); a.fill.fore_color.rgb = col; a.line.fill.background()

def arrow(s, x1, y1, x2, y2):
    cn = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    cn.line.color.rgb = BLUE; cn.line.width = Pt(3.5)
    ln = cn.line._get_or_add_ln()
    etree.SubElement(ln, qn('a:tailEnd'), type='triangle', w='lg', len='lg')


# ============================ SLIDE 1: Strategy & Overall ======================
s1 = new_slide('Testing Strategy & Overall Results')
kp1 = [
    ('352', 'TOTAL TESTS', BLUE),
    ('302', 'AUTOMATED UNIT TESTS', BLUE),
    ('50', 'SYSTEM TEST CASES', GREEN),
    ('0', 'DEFECTS REMAINING', GREEN),
    ('100%', 'OVERALL PASS RATE', GREEN)
]
gap = 0.22; kw = (18.0 - gap * 4) / 5
for i, (v, lab, col) in enumerate(kp1):
    x = 1.0 + i * (kw + gap)
    box(s1, x, 1.75, kw, 1.05, radius=0.12); accent(s1, x, 1.75, kw, col)
    text(s1, x, 1.87, kw, 0.55, v, 28, col, True)
    text(s1, x, 2.43, kw, 0.28, lab, 11, GREY, True)

# Left Column: Tier 1 Automated Unit Testing
cw, ch = 8.8, 7.1
box(s1, 1.0, 3.1, cw, ch); accent(s1, 1.0, 3.1, cw, BLUE)
text(s1, 1.3, 3.3, cw - 0.6, 0.45, 'Tier 1: Automated Unit Testing (Code Level)', 18, DBLUE, True, PP_ALIGN.LEFT)
text(s1, 1.3, 3.75, cw - 0.6, 0.35, '302 Automated Tests · 24 Test Suites · 100% Pass Rate', 13, BLUE, True, PP_ALIGN.LEFT)

# Sub-card 1: Backend
box(s1, 1.3, 4.25, cw - 0.6, 2.3, fill=ROWALT, line=LIGHT, radius=0.06)
text(s1, 1.5, 4.35, cw - 1.0, 0.35, 'Backend Application Server (NestJS + Jest)', 15, DBLUE, True, PP_ALIGN.LEFT)
b_pts = [
    '• 21 Test Suites · 257 Automated Tests (Executed in 23.4s)',
    '• Authentication, JWT token lifecycle, password & email hashing',
    '• Role-based access control (Admin, Tester, Viewer) & project security',
    '• OpenAPI 3.0 specification parsing, endpoint extraction, report exports'
]
for j, pt in enumerate(b_pts):
    text(s1, 1.5, 4.75 + j * 0.42, cw - 1.0, 0.38, pt, 12, GREY if j > 0 else DBLUE, j == 0, PP_ALIGN.LEFT)

# Sub-card 2: Engine
box(s1, 1.3, 6.75, cw - 0.6, 2.3, fill=ROWALT, line=LIGHT, radius=0.06)
text(s1, 1.5, 6.85, cw - 1.0, 0.35, 'AI Engine Microservice (Python + Pytest)', 15, DBLUE, True, PP_ALIGN.LEFT)
e_pts = [
    '• 3 Test Modules · 45 Automated Tests (Executed in 13.9s)',
    '• Engine REST API: Run dispatching, bearer tokens, status polling',
    '• OOPS Code-to-Spec: Source zip parsing, file traversal security',
    '• Execution Runner: RL process lifecycle, proxy request capture'
]
for j, pt in enumerate(e_pts):
    text(s1, 1.5, 7.25 + j * 0.42, cw - 1.0, 0.38, pt, 12, GREY if j > 0 else DBLUE, j == 0, PP_ALIGN.LEFT)

# Bottom pill left
box(s1, 1.3, 9.25, cw - 0.6, 0.65, fill=GREENBG, line=None, radius=0.5)
text(s1, 1.3, 9.25, cw - 0.6, 0.65, '✔ Continuous Integration: 302/302 Passed (~37s total run)', 13, GREEN, True)


# Right Column: Tier 2 System & Black-Box Testing
box(s1, 10.2, 3.1, cw, ch); accent(s1, 10.2, 3.1, cw, ORANGE)
text(s1, 10.5, 3.3, cw - 0.6, 0.45, 'Tier 2: System & Black-Box Testing (E2E)', 18, DBLUE, True, PP_ALIGN.LEFT)
text(s1, 10.5, 3.75, cw - 0.6, 0.35, '50 End-to-End Scenarios · 7 Functional Modules · 100% Pass', 13, ORANGE, True, PP_ALIGN.LEFT)

# Sub-card 1: Deployed Platform
box(s1, 10.5, 4.25, cw - 0.6, 2.3, fill=ROWALT, line=LIGHT, radius=0.06)
text(s1, 10.7, 4.35, cw - 1.0, 0.35, 'Deployed Testing Environment', 15, DBLUE, True, PP_ALIGN.LEFT)
s_pts = [
    '• Fully Deployed SaaS Stack (Vercel + Render + NeonDB)',
    '• Exercised across both web UI and underlying REST API endpoints',
    '• Multi-account concurrent testing (Owner, Admin, Tester, Viewer)',
    '• Live AI provider interaction + mock deterministic mode'
]
for j, pt in enumerate(s_pts):
    text(s1, 10.7, 4.75 + j * 0.42, cw - 1.0, 0.38, pt, 12, GREY if j > 0 else DBLUE, j == 0, PP_ALIGN.LEFT)

# Sub-card 2: Categorization
box(s1, 10.5, 6.75, cw - 0.6, 2.3, fill=ROWALT, line=LIGHT, radius=0.06)
text(s1, 10.7, 6.85, cw - 1.0, 0.35, 'Test Case Intent & Breakdown', 15, DBLUE, True, PP_ALIGN.LEFT)
t_pts = [
    '• Functional Workflows: 35 Cases (Happy-path operational flows)',
    '• Negative / Validation: 10 Cases (Bad inputs, malformed specs, sizes)',
    '• Security & Permissions: 5 Cases (RBAC rules, unverified states)',
    '• Modules Covered: Auth, Projects, Specs, Suites, Execution, Reports, Team'
]
for j, pt in enumerate(t_pts):
    text(s1, 10.7, 7.25 + j * 0.42, cw - 1.0, 0.38, pt, 12, GREY if j > 0 else DBLUE, j == 0, PP_ALIGN.LEFT)

# Bottom pill right
box(s1, 10.5, 9.25, cw - 0.6, 0.65, fill=GREENBG, line=None, radius=0.5)
text(s1, 10.5, 9.25, cw - 0.6, 0.65, '✔ Full Acceptance: 50/50 Passed with 0 Failures', 13, GREEN, True)


# ============================ SLIDE 2: Automated Unit Testing ==================
s2 = new_slide('Automated Unit Testing')
kp2 = [
    ('302', 'TOTAL UNIT TESTS', BLUE),
    ('24', 'TEST SUITES', BLUE),
    ('257', 'NESTJS (JEST)', BLUE),
    ('45', 'PYTHON (PYTEST)', ORANGE),
    ('37.3s', 'CI RUNTIME', GREEN)
]
for i, (v, lab, col) in enumerate(kp2):
    x = 1.0 + i * (kw + gap)
    box(s2, x, 1.75, kw, 1.05, radius=0.12); accent(s2, x, 1.75, kw, col)
    text(s2, x, 1.87, kw, 0.55, v, 28, col, True)
    text(s2, x, 2.43, kw, 0.28, lab, 11, GREY, True)

# Left Column: Backend (NestJS + Jest)
box(s2, 1.0, 3.1, cw, 6.0); accent(s2, 1.0, 3.1, cw, BLUE)
text(s2, 1.3, 3.3, cw - 0.6, 0.45, 'Backend Application Server (NestJS + Jest)', 18, DBLUE, True, PP_ALIGN.LEFT)
text(s2, 1.3, 3.75, cw - 0.6, 0.35, '21 Test Suites · 257 Tests · 23.39s Runtime · 100% Pass Rate', 13, BLUE, True, PP_ALIGN.LEFT)

b_domains = [
    ('Authentication & Token Security (41 tests)', 'auth.service, signup.service, tokens.service, jwt.strategy', 'Password hashing, JWT cookie signing, 6-digit confirmation logic'),
    ('RBAC & Project Isolation (40 tests)', 'project-access.service, members.service, invitations.service', 'Admin/Tester/Viewer boundaries, multi-tenant data guards'),
    ('OpenAPI & Spec Handling (42 tests)', 'endpoint-extractor, specs.service, spec-generation.service', 'OAS 3.0 parsing, verb extraction, zip archive security validation'),
    ('Engines, Reports & Utilities (134 tests)', 'engine.service, reports.service, report-export, graph-merge, ...', 'SPDG graph union, PDF/CSV export, LLM prompt formatting, email mock')
]

for d_i, (d_title, d_files, d_desc) in enumerate(b_domains):
    dy = 4.25 + d_i * 1.15
    box(s2, 1.3, dy, cw - 0.6, 1.02, fill=ROWALT, line=LIGHT, radius=0.06)
    text(s2, 1.5, dy + 0.1, cw - 1.0, 0.3, d_title, 14, DBLUE, True, PP_ALIGN.LEFT)
    text(s2, 1.5, dy + 0.42, cw - 1.0, 0.26, 'Files: ' + d_files, 11, BLUE, False, PP_ALIGN.LEFT)
    text(s2, 1.5, dy + 0.68, cw - 1.0, 0.26, d_desc, 11, GREY, False, PP_ALIGN.LEFT)

# Right Column: Engine (Python + Pytest)
box(s2, 10.2, 3.1, cw, 6.0); accent(s2, 10.2, 3.1, cw, ORANGE)
text(s2, 10.5, 3.3, cw - 0.6, 0.45, 'AI Engine Microservice (Python + Pytest)', 18, DBLUE, True, PP_ALIGN.LEFT)
text(s2, 10.5, 3.75, cw - 0.6, 0.35, '3 Test Modules · 45 Tests · 13.96s Runtime · 100% Pass Rate', 13, ORANGE, True, PP_ALIGN.LEFT)

e_domains = [
    ('Engine REST API Service (8 tests)', 'tests/test_api.py', 'Microservice endpoints, health probes, bearer authentication, run job creation, status polling, and error status code translation.'),
    ('OOPS Spec Generation Pipeline (18 tests)', 'tests/test_generation.py', 'Source code zip extraction, file filtering, path traversal protection (../ attacks), rate limits, and LLM request timeouts.'),
    ('MARL Runner & Proxy Engine (19 tests)', 'tests/test_runner.py', 'Multi-agent RL testing runner process lifecycle, proxy request capture and interception, cached values, and deterministic mock execution.')
]

for d_i, (d_title, d_files, d_desc) in enumerate(e_domains):
    dy = 4.25 + d_i * 1.55
    box(s2, 10.5, dy, cw - 0.6, 1.4, fill=ROWALT, line=LIGHT, radius=0.06)
    text(s2, 10.7, dy + 0.12, cw - 1.0, 0.32, d_title, 14, DBLUE, True, PP_ALIGN.LEFT)
    text(s2, 10.7, dy + 0.46, cw - 1.0, 0.28, 'File: ' + d_files, 11, ORANGE, False, PP_ALIGN.LEFT)
    text(s2, 10.7, dy + 0.78, cw - 1.0, 0.52, d_desc, 11, GREY, False, PP_ALIGN.LEFT)

# Bottom Integration Card
box(s2, 1.0, 9.3, 18.0, 0.8, fill=WHITE, line=LIGHT, radius=0.08); accent(s2, 1.0, 9.3, 18.0, GREEN)
text(s2, 1.3, 9.42, 17.4, 0.55, 'Continuous Integration & Quality Assurance: All 302 unit tests run automatically in CI under 40 seconds with 100% repeatability, ensuring no architectural regression reaches production.', 13, DBLUE, True, PP_ALIGN.CENTER)


# ============================ SLIDE 3: Module-wise table =======================
s3 = new_slide('System Testing: Module-wise Coverage')
cols = ['Module', 'Cases', 'Result', 'What was verified']
cwi = [4.6, 1.4, 1.8, 10.2]
y0 = 1.8; hh = 0.55; rh = 0.98
gf = s3.shapes.add_table(len(MODULES) + 2, 4, E(1.0), E(y0), E(sum(cwi)), E(hh + rh * len(MODULES) + 0.7))
tbl = gf.table
tbl._tbl.tblPr.find(qn('a:tableStyleId')).text = '{2D5ABB26-0587-4C30-8999-92F81FD0307C}'
for i, w in enumerate(cwi):
    tbl.columns[i].width = E(w)
tbl.rows[0].height = E(hh)
for i in range(1, len(MODULES) + 1):
    tbl.rows[i].height = E(rh)
tbl.rows[len(MODULES) + 1].height = E(0.7)

def cell(r, c, txt, size, color, bold, align, fill):
    ce = tbl.cell(r, c)
    ce.margin_left = ce.margin_right = E(0.18); ce.margin_top = ce.margin_bottom = 0
    ce.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = ce.text_frame.paragraphs[0]; p.alignment = align
    ru = p.add_run(); ru.text = txt
    ru.font.size = Pt(size); ru.font.bold = bold; ru.font.color.rgb = color; ru.font.name = FONT
    ce.fill.solid(); ce.fill.fore_color.rgb = fill

for c, h in enumerate(cols):
    cell(0, c, h, 14, WHITE, True, PP_ALIGN.LEFT if c in (0, 3) else PP_ALIGN.CENTER, DBLUE)
for i, (nm, n, ver) in enumerate(MODULES, start=1):
    bg = ROWALT if i % 2 == 0 else WHITE
    cell(i, 0, f'{i}.  {nm}', 17, DBLUE, True, PP_ALIGN.LEFT, bg)
    cell(i, 1, str(n), 20, MCOL[i - 1], True, PP_ALIGN.CENTER, bg)
    cell(i, 2, 'PASS', 15, GREEN, True, PP_ALIGN.CENTER, GREENBG)
    cell(i, 3, ver, 15, GREY, False, PP_ALIGN.LEFT, bg)
tr = len(MODULES) + 1
cell(tr, 0, 'Total', 17, WHITE, True, PP_ALIGN.LEFT, BLUE)
cell(tr, 1, '50', 20, WHITE, True, PP_ALIGN.CENTER, BLUE)
cell(tr, 2, '50 / 50', 15, WHITE, True, PP_ALIGN.CENTER, GREEN)
cell(tr, 3, '0 failures across all seven modules', 15, WHITE, True, PP_ALIGN.LEFT, BLUE)


# ============================ SLIDE 4: What tests proved =======================
s4 = new_slide('What the Tests Proved')
cards = [
    ('1', 'Input validation', BLUE, 'Invalid or missing fields are rejected with field-specific messages; no partial records are created.', 'TC 2, 3, 6, 18, 31'),
    ('2', 'Security', ORANGE, 'Confirmation codes are stored hashed, with no session before email confirmation. Path-traversal and symlink uploads are refused.', 'TC 1, 22, 46'),
    ('3', 'Access control', GREEN, 'Viewer, tester and admin boundaries hold; only the owner can delete a project; removed members lose access at once.', 'TC 15, 20, 48, 49'),
    ('4', 'Robustness', PURPLE, 'No run without a specification, duplicate runs refused, asynchronous runs keep the UI responsive.', 'TC 33, 34, 35'),
]
cw3, ch3 = 8.8, 2.55
for i, (n, hd, col, body, ref) in enumerate(cards):
    x = 1.0 + (i % 2) * (cw3 + 0.4); y = 1.8 + (i // 2) * (ch3 + 0.3)
    box(s4, x, y, cw3, ch3); accent(s4, x, y, cw3, col)
    box(s4, x + 0.3, y + 0.35, 0.7, 0.7, col, None, shape=MSO_SHAPE.OVAL)
    text(s4, x + 0.3, y + 0.35, 0.7, 0.7, n, 20, WHITE, True)
    text(s4, x + 1.25, y + 0.35, 5.0, 0.7, hd, 21, DBLUE, True, PP_ALIGN.LEFT)
    text(s4, x + 6.3, y + 0.45, 2.3, 0.5, ref, 12, col, True, PP_ALIGN.RIGHT)
    text(s4, x + 0.3, y + 1.2, cw3 - 0.6, 1.2, body, 17, GREY, False, PP_ALIGN.LEFT, MSO_ANCHOR.TOP)

# worked example strip
ey = 7.45
box(s4, 1.0, ey, 18.0, 3.0, LIGHT, None, 0.06)
text(s4, 1.3, ey + 0.15, 9.0, 0.45, 'Worked example  |  TC33 One-Click Test Execution', 16, DBLUE, True, PP_ALIGN.LEFT)
steps = [('STEPS', BLUE, 'Press Run, watch the status without reloading, navigate away and come back'),
         ('EXPECTED', ORANGE, 'Run returns instantly; status goes pending → running → completed; UI stays usable'),
         ('RESULT', GREEN, 'PASS  –  state is current on return, never stale')]
bw = 5.3; bx0 = 1.3; gp = 0.9
for i, (hd, col, body) in enumerate(steps):
    x = bx0 + i * (bw + gp)
    box(s4, x, ey + 0.75, bw, 1.95); accent(s4, x, ey + 0.75, bw, col)
    text(s4, x, ey + 0.9, bw, 0.4, hd, 14, col, True)
    text(s4, x + 0.3, ey + 1.35, bw - 0.6, 1.25, body, 14, GREY, False, PP_ALIGN.CENTER)
    if i < 2:
        arrow(s4, x + bw + 0.1, ey + 1.72, x + bw + gp - 0.1, ey + 1.72)


# ============================ SLIDE 5: Appendix all 50 =========================
s5 = new_slide('Appendix: All 50 Test Cases')
idx = 0; mod_of = []
for mi, m in enumerate(MODULES):
    mod_of += [mi] * m[1]
rows = 25; rowh = 0.355; ay = 1.75
colw = 8.8
for k in range(50):
    c = k // rows; r = k % rows
    x = 1.0 + c * (colw + 0.4); y = ay + r * rowh
    mc = MCOL[mod_of[k]]
    box(s5, x, y + 0.03, 0.62, rowh - 0.06, mc, None, 0.3)
    text(s5, x, y + 0.03, 0.62, rowh - 0.06, str(k + 1), 11, WHITE, True)
    text(s5, x + 0.8, y, colw - 1.6, rowh, TITLES[k], 12, GREY, False, PP_ALIGN.LEFT)
    text(s5, x + colw - 0.7, y, 0.7, rowh, 'Pass', 11, GREEN, True, PP_ALIGN.RIGHT)

lx = 1.0; ly = ay + rows * rowh + 0.2
for i, m in enumerate(MODULES):
    box(s5, lx, ly + 0.07, 0.3, 0.2, MCOL[i], None, 0.5)
    text(s5, lx + 0.4, ly, 2.2, 0.35, m[0], 11, GREY, False, PP_ALIGN.LEFT)
    lx += 2.6

first = list(ids)[0]
prs.part.drop_rel(first.get(qn('r:id'))); ids.remove(first)
prs.save(OUT)
print('saved', OUT, len(prs.slides))
