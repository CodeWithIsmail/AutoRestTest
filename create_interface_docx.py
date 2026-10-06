import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x1F, 0x39, 0x64) # Navy blue
        if level == 2:
            run.font.size = Pt(14)
            run.bold = True
        elif level == 3:
            run.font.size = Pt(13)
            run.bold = True
        elif level == 4:
            run.font.size = Pt(12)
            run.bold = True
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.space_before = Pt(12)

def add_paragraph(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p

def set_table_header(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        # Add light gray background to header
        shading_elm = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading_elm)
        
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.bold = True

def add_table_row(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)

def main():
    doc = Document()

    # Match the user's exact margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    add_heading(doc, '3.2 Interface Design', 2)
    
    add_heading(doc, '3.2.1 User Interface Architecture', 3)
    add_paragraph(doc, "The frontend utilizes a highly modular, component-based architecture built with React and Next.js, with React Query (TanStack Query) orchestrating asynchronous data fetching and global state management. The user interface abstracts the complexity of the underlying LLM testing engine, providing dedicated objects for managing AI execution flows. Below are the core UI objects and their hierarchical relationships:")
    
    add_paragraph(doc, "Table 3.1: UI Objects and Actions", italic=True)

    # Table 1: UI Objects and Actions
    table1 = doc.add_table(rows=1, cols=3)
    table1.style = 'Table Grid'
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    widths1 = [Inches(1.8), Inches(1.5), Inches(3.2)]
    for i in range(3):
        table1.columns[i].width = widths1[i]
        for cell in table1.columns[i].cells:
            cell.width = widths1[i]

    set_table_header(table1.rows[0], ['UI Object', 'Type', 'Actions (Operations)'])

    data1 = [
        ("SpecUploader", "Dropzone Form", "Upload OpenAPI schema (.yaml/.json), triggering parameter extraction"),
        ("SpecGenerator", "Multi-step Form", "Upload source code zip, poll AI parsing status, apply generated schema"),
        ("EndpointsList", "Data Display", "Filter parsed API routes, manually configure parameters for testing"),
        ("DependencyGraph", "Interactive Canvas", "Render endpoints as connected nodes, pan/zoom, trigger embeddings build"),
        ("FocusView", "Canvas Extension", "Isolate parent/child graph relations in a rigid 3-column topology"),
        ("TestSuiteBuilder", "Form", "Configure target URL/time budgets, dispatch test generation tasks to MARL engine"),
        ("TestTimeline", "Live Display", "Monitor execution progress, render live request status and pass/fail metrics"),
        ("LogViewer", "Accordion", "Expand historical HTTP records, inspect raw JSON request/response payloads"),
        ("LiveReplayer", "Control", "Re-execute a sequence of API calls to verify stability or code fixes"),
        ("AIExplanationModal", "Floating Widget", "Stream LLM-generated root-cause analyses for test failures"),
        ("ReportExporter", "Control", "Download test execution summaries as PDF or CSV artifacts"),
        ("LLMSettings", "Admin Panel", "Live-tune NVIDIA API keys and foundation models without redeploying")
    ]

    for obj, typ, actions in data1:
        row = table1.add_row()
        add_table_row(row, [obj, typ, actions])

    add_paragraph(doc, "")
    
    add_heading(doc, '3.2.2 UI Events and State Transitions', 3)
    
    add_paragraph(doc, "Table 3.2: Key UI Events and State Changes", italic=True)

    # Table 2: UI Events and State Transitions
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = 'Table Grid'
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    widths2 = [Inches(1.8), Inches(2.2), Inches(2.5)]
    for i in range(3):
        table2.columns[i].width = widths2[i]
        for cell in table2.columns[i].cells:
            cell.width = widths2[i]

    set_table_header(table2.rows[0], ['User Event', 'Action', 'UI State Change'])

    data2 = [
        ("Upload Spec File", "POST /projects/:id/spec", "Editor mounts YAML, API extracts routes into EndpointsList"),
        ("Generate Spec (AI)", "POST .../spec/generate", "Polling state activates -> AI progress bar tracks completion"),
        ("Build Dependency Graph", "POST /projects/:id/graph", "Polling begins -> Canvas re-renders when graph embeddings resolve"),
        ("Toggle Focus Mode", "Local State Update", "Canvas collapses non-related nodes, strictly sorting parents/children"),
        ("Create Test Run", "POST /projects/:id/test-suites", "Dispatches to MARL engine -> UI creates pending test-suite row"),
        ("Execute Test Suite", "POST .../test-suites/:id/run", "Status turns 'Running', TestTimeline streams progress dynamically"),
        ("Expand Request Log", "GET .../test-suites/:id/test-cases", "Row expands, HTTP headers and JSON payloads mount"),
        ("Click \"AI Explain\"", "POST .../test-suites/:id/explain", "Modal opens, text streams detailing failure root cause"),
        ("Replay Execution", "POST .../test-suites/:id/replay", "Dispatches replay task, UI updates with new response data"),
        ("Export Run Report", "GET .../test-suites/:id/report/export", "Blob streams to client, browser triggers file download"),
        ("Update LLM Config", "PATCH /admin/llm-settings/:scope", "Configuration hot-reloads, subsequent AI calls use new parameters")
    ]

    for event, action, state_change in data2:
        row = table2.add_row()
        add_table_row(row, [event, action, state_change])

    add_paragraph(doc, "")

    add_heading(doc, '3.2.3 Interface States Depiction', 3)
    add_paragraph(doc, "The following UI depictions illustrate the actual interface states as they appear to the end user across the core testing workflows:")

    add_heading(doc, 'Project Initialization & Graph Architecture', 4)
    add_paragraph(doc, "• Project Overview: The high-level workspace summary. (Insert: '9. project overview.png')")
    add_paragraph(doc, "• Specification Management: The state showing uploaded or AI-generated OpenAPI schemas. (Insert: '10.1 api spec.png', '10.2 generate api spec.png')")
    add_paragraph(doc, "• Endpoints Configuration: The listing interface for API routes. (Insert: '11.1. endpoints.png', '11.2 add endpoint.png')")
    add_paragraph(doc, "• Dependency Graph: The interactive visualization of API execution order, utilizing the Focus View. (Insert: '12. graph.png')")

    add_heading(doc, 'Test Execution & Live Replay', 4)
    add_paragraph(doc, "• Suite Orchestration: The configuration state for dispatching test cases. (Insert: '13. test suit create.png')")
    add_paragraph(doc, "• Execution Timeline & Results: The detailed reporting state displaying granular HTTP request payloads, live request replays, and AI Explanations. (Insert: '14.1 test result.png', '14.2 request log.png')")

    add_heading(doc, 'Global Navigation & Settings', 4)
    add_paragraph(doc, "• Global Dashboard: The central hub displaying active projects. (Insert: '7. homepage dashboard.png')")
    add_paragraph(doc, "• LLM Administration: The state for configuring foundational models and API keys. (Insert: '16. profile settings.png')")

    output_path_root = os.path.join("D:\\SPL3\\AutoRestTest", "AutoRestTest_InterfaceDesign_Final_v3.docx")
    output_path_srs = os.path.join("D:\\SPL3\\AutoRestTest", "resources", "SRS", "AutoRestTest_InterfaceDesign_Final_v3.docx")
    
    doc.save(output_path_root)
    os.makedirs(os.path.dirname(output_path_srs), exist_ok=True)
    doc.save(output_path_srs)

if __name__ == "__main__":
    main()
