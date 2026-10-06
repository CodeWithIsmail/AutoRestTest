import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_document():
    doc = docx.Document()
    
    # Page setup matching SPL3 Technical Report (Letter, 1" top/bottom, 1.25" left/right)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    
    # Configure Normal Style
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(4)
    normal_style.paragraph_format.space_before = Pt(0)
    
    NAVY = RGBColor(0x1F, 0x39, 0x64)
    BLACK = RGBColor(0x00, 0x00, 0x00)
    
    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_h4(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.bold = True
        run.font.color.rgb = NAVY
        return p

    def add_req(code_and_title, description):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r1 = p.add_run(code_and_title + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(12)
        r1.bold = True
        r1.font.color.rgb = BLACK
        
        r2 = p.add_run(description)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(12)
        r2.bold = False
        r2.font.color.rgb = BLACK
        return p

    def add_italic_note(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        run.italic = True
        run.font.color.rgb = BLACK
        return p

    # --- Document Content ---
    add_h2("2.1 Quality Function Deployment")
    
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(2)
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.15
    r_intro = p_intro.add_run(
        "Quality Function Deployment (QFD), combined with the Kano model of customer satisfaction, "
        "categorizes the platform's requirements into three distinct tiers: Normal Requirements "
        "(explicitly requested features where user satisfaction scales proportionally with performance), "
        "Expected Requirements (must-have baseline features whose absence causes severe dissatisfaction, "
        "even if users take them for granted), and Exciting Requirements (innovative delighters that "
        "differentiate the platform and exceed user expectations)."
    )
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(12)

    # -------------------------------------------------------------
    # 2.1.1 Normal Requirements
    # -------------------------------------------------------------
    add_h3("2.1.1 Normal Requirements")
    add_italic_note("Explicitly requested core features; user satisfaction is proportional to how well they are met.")
    
    add_h4("Functional Requirements")
    add_req("FR-N1 \u2014 User Authentication & Session Management", 
            "register new accounts, authenticate using email or username, and maintain secure, stateless JWT-based sessions.")
    add_req("FR-N2 \u2014 Project Management", 
            "create, open, rename, search, filter (by user role and specification status), and delete API testing workspaces from a central dashboard with card and table layouts.")
    add_req("FR-N3 \u2014 Specification Upload & Ingestion", 
            "upload OpenAPI 3.0 or Swagger specifications in YAML or JSON format; the system parses, validates, and stores the specification while allowing users to inspect, copy, or download the raw content.")
    add_req("FR-N4 \u2014 API Endpoint Exploration & Manual Management", 
            "automatically extract and list all endpoints, operations, parameter schemas (path, query, header, cookie), and response models from the specification, with support for manually adding or deleting custom endpoints.")
    add_req("FR-N5 \u2014 Configurable Automated Test Execution", 
            "configure test runs against a live target base URL with customizable execution time budgets (in seconds), mutation fuzzing rates (0.0 to 1.0), custom request headers (up to 20 headers, e.g., Bearer auth or API tokens), and operation exclusion filters to bypass flaky or destructive endpoints.")
    add_req("FR-N6 \u2014 Granular Result & Request Log Inspection", 
            "view run execution outcomes per endpoint and browse the comprehensive sequence of individual HTTP requests and responses captured by the recording proxy (including HTTP method, concrete path, status code, duration, full request/response headers, and payloads).")
    add_req("FR-N7 \u2014 Completion & System Notifications", 
            "deliver real-time in-app status updates and transactional email alerts to notify users when an asynchronous test execution completes or encounters fatal errors.")

    add_h4("Non-Functional Requirements")
    add_req("NFR-N1 \u2014 Usability & UI Ergonomics", 
            "the whole workflow is driven from one intuitive, modern web dashboard featuring top-bar navigation, a quick project switcher, notification badges, and a light/dark theme toggle, usable without special training.")
    add_req("NFR-N2 \u2014 Interactive System Responsiveness", 
            "routine interactions (dashboard filtering, sorting, endpoint exploration, and paginated log retrieval) execute and render within a couple of seconds.")

    # -------------------------------------------------------------
    # 2.1.2 Expected Requirements
    # -------------------------------------------------------------
    add_h3("2.1.2 Expected Requirements")
    add_italic_note("Implicitly assumed; their absence causes dissatisfaction even though users rarely ask for them by name.")

    add_h4("Functional Requirements")
    add_req("FR-E1 \u2014 Automatic Test Suite Generation", 
            "automatically generate a comprehensive, dependency-aware suite of test cases from the specification, resolving inter-operation parameter relationships without manual test scripting.")
    add_req("FR-E2 \u2014 Deterministic Regression Test Replay", 
            "resend the exact sequence of captured requests from a completed test run verbatim against the target API (as a linked replay run) to verify defect resolutions deterministically without stochastic AI variation.")
    add_req("FR-E3 \u2014 Execution History & Replay Lineage Tracking", 
            "persist complete test execution histories with execution timestamps, durations, and outcomes, maintaining explicit lineage links between original AI-driven baseline runs and subsequent replay regression runs.")
    add_req("FR-E4 \u2014 Reports, Analytics & Fault Discrimination", 
            "per-run summary of endpoint coverage percentage, pass/fail rates, status-code distribution, and clear visual discrimination between expected client rejections (4xx validation responses) and genuine server failures (5xx crashes).")
    add_req("FR-E5 \u2014 Export Reports", 
            "download a completed run's comprehensive report and results breakdown in standard formats such as PDF or CSV.")
    add_req("FR-E6 \u2014 Team Collaboration & Role-Based Access Control", 
            "invite collaborators to a project via email with tokenized invitations, accept/decline invites via a centralized invitations hub, and enforce granular role-based permissions (Owner, Admin, Tester, Viewer) across specifications, runs, and settings.")
    add_req("FR-E7 \u2014 Account Security & Profile Lifecycle", 
            "self-service six-digit email verification with anti-squatting protection, token-based password reset, active session revocation upon credential change, profile customization, and password-confirmed account deletion.")

    add_h4("Non-Functional Requirements")
    add_req("NFR-E1 \u2014 Security & Data Privacy", 
            "passwords hashed with bcrypt, single-use authentication tokens hashed with SHA-256, private routes protected by JWT guards, and sensitive execution credentials (e.g., custom Authorization headers, LLM API keys) masked as write-only fields.")
    add_req("NFR-E2 \u2014 Reliability & Asynchronous Execution", 
            "long-running test executions, codebase analysis jobs, and dependency graph builds run asynchronously in decoupled background workers without blocking the UI, reporting real-time status accurately.")
    add_req("NFR-E3 \u2014 Data Persistence & Integrity", 
            "specifications, suites, runs, captured request/response pairs, and team permissions are stored durably in PostgreSQL via Prisma ORM and remain available across sessions.")
    add_req("NFR-E4 \u2014 Brute-Force & Abuse Resistance", 
            "verification codes and password-reset tokens enforce strict attempt ceilings (auto-lockout on repeated failures) and IP-based rate limiting to protect authentication endpoints from brute-force attacks.")

    # -------------------------------------------------------------
    # 2.1.3 Exciting Requirements
    # -------------------------------------------------------------
    add_h3("2.1.3 Exciting Requirements")
    add_italic_note("Delighters beyond expectation; their absence isn't penalised, but their presence differentiates the platform.")

    add_h4("Functional Requirements")
    add_req("FR-X1 \u2014 One-Click End-to-End Testing", 
            "a single action triggers the autonomous testing pipeline\u2014reading the active specification, synthesizing dependency-aware test cases, executing requests against the live target, and producing an analytics-rich report without manual intervention.")
    add_req("FR-X2 \u2014 AI Specification Generation from Source Code", 
            "automatically generate a valid OpenAPI 3.0 specification from an uploaded codebase archive (.zip) via a 9-step AST and LLM extraction pipeline, complete with directory exclusion, warning detection, and a review-before-apply diff stage.")
    add_req("FR-X3 \u2014 Interactive Semantic Dependency Graph Visualization", 
            "build and visualize the Semantic Operation Dependency Graph (SODG) using an interactive 3-column neighborhood layout (Producers \u2192 Focused Operation \u2192 Consumers) displaying semantic edge strengths, cycle counts, entry points, and run-scoped snapshots enriched with learned Multi-Agent Reinforcement Learning (MARL) Q-values.")
    add_req("FR-X4 \u2014 AI Failure Explanation", 
            "LLM-generated, plain-language diagnostic explanation of what caused a server failure (5xx), which specific request and payload triggered it, and recommended remediation guidance.")
    add_req("FR-X5 \u2014 AI Request Test Intent Explanation", 
            "provide on-demand batch LLM summarization that annotates individual captured HTTP requests with concise, plain-language descriptions explaining what specific boundary condition, mutation, or input property is being exercised.")
    add_req("FR-X6 \u2014 Interactive Live Request Re-Execution & cURL Generation", 
            "resend any individual captured request directly to the target API on demand from the web interface to observe live response changes, with safety confirmations for state-mutating HTTP methods and one-click 'Copy as cURL' generation.")
    add_req("FR-X7 \u2014 Multi-Scope Centralized LLM Orchestration", 
            "provide an administrative configuration interface allowing operators to independently configure AI models, custom API base URLs, provider keys, and Requests-Per-Minute (RPM) rate limits across three decoupled AI functional scopes: Test Engine, Spec Generation, and Report Explanation.")

    add_h4("Non-Functional Requirements")
    add_req("NFR-X1 \u2014 Modular Microservice & Engine Extensibility", 
            "the AI testing engine and spec generation engine are deployed as independent headless services communicating over internal REST APIs, allowing generation strategies, Multi-Agent Reinforcement Learning (MARL) algorithms, and LLM providers to evolve without changing the web platform.")
    add_req("NFR-X2 \u2014 Scalable Graph Rendering Performance", 
            "the dependency graph visualizer utilizes localized O(1) neighborhood projection rather than full monolithic graph rendering, ensuring smooth, lag-free UI interaction even on enterprise APIs with hundreds of interdependent endpoints.")

    # -------------------------------------------------------------
    # Summary Table
    # -------------------------------------------------------------
    p_tbl_title = doc.add_paragraph()
    p_tbl_title.paragraph_format.space_before = Pt(14)
    p_tbl_title.paragraph_format.space_after = Pt(4)
    p_tbl_title.paragraph_format.keep_with_next = True
    r_tbl_title = p_tbl_title.add_run("Table 2.1: Quality Function Deployment (QFD) Requirements Summary")
    r_tbl_title.font.name = 'Times New Roman'
    r_tbl_title.font.size = Pt(12)
    r_tbl_title.bold = True
    r_tbl_title.font.color.rgb = NAVY

    table_data = [
        ("Kano Category", "Requirement ID", "Requirement Title", "Type"),
        # Normal
        ("Normal Requirements", "FR-N1", "User Authentication & Session Management", "Functional"),
        ("Normal Requirements", "FR-N2", "Project Management", "Functional"),
        ("Normal Requirements", "FR-N3", "Specification Upload & Ingestion", "Functional"),
        ("Normal Requirements", "FR-N4", "API Endpoint Exploration & Manual Management", "Functional"),
        ("Normal Requirements", "FR-N5", "Configurable Automated Test Execution", "Functional"),
        ("Normal Requirements", "FR-N6", "Granular Result & Request Log Inspection", "Functional"),
        ("Normal Requirements", "FR-N7", "Completion & System Notifications", "Functional"),
        ("Normal Requirements", "NFR-N1", "Usability & UI Ergonomics", "Non-Functional"),
        ("Normal Requirements", "NFR-N2", "Interactive System Responsiveness", "Non-Functional"),
        # Expected
        ("Expected Requirements", "FR-E1", "Automatic Test Suite Generation", "Functional"),
        ("Expected Requirements", "FR-E2", "Deterministic Regression Test Replay", "Functional"),
        ("Expected Requirements", "FR-E3", "Execution History & Replay Lineage Tracking", "Functional"),
        ("Expected Requirements", "FR-E4", "Reports, Analytics & Fault Discrimination", "Functional"),
        ("Expected Requirements", "FR-E5", "Export Reports", "Functional"),
        ("Expected Requirements", "FR-E6", "Team Collaboration & Role-Based Access Control", "Functional"),
        ("Expected Requirements", "FR-E7", "Account Security & Profile Lifecycle", "Functional"),
        ("Expected Requirements", "NFR-E1", "Security & Data Privacy", "Non-Functional"),
        ("Expected Requirements", "NFR-E2", "Reliability & Asynchronous Execution", "Non-Functional"),
        ("Expected Requirements", "NFR-E3", "Data Persistence & Integrity", "Non-Functional"),
        ("Expected Requirements", "NFR-E4", "Brute-Force & Abuse Resistance", "Non-Functional"),
        # Exciting
        ("Exciting Requirements", "FR-X1", "One-Click End-to-End Testing", "Functional"),
        ("Exciting Requirements", "FR-X2", "AI Specification Generation from Source Code", "Functional"),
        ("Exciting Requirements", "FR-X3", "Interactive Semantic Dependency Graph Visualization", "Functional"),
        ("Exciting Requirements", "FR-X4", "AI Failure Explanation", "Functional"),
        ("Exciting Requirements", "FR-X5", "AI Request Test Intent Explanation", "Functional"),
        ("Exciting Requirements", "FR-X6", "Interactive Live Request Re-Execution & cURL Generation", "Functional"),
        ("Exciting Requirements", "FR-X7", "Multi-Scope Centralized LLM Orchestration", "Functional"),
        ("Exciting Requirements", "NFR-X1", "Modular Microservice & Engine Extensibility", "Non-Functional"),
        ("Exciting Requirements", "NFR-X2", "Scalable Graph Rendering Performance", "Non-Functional"),
    ]

    table = doc.add_table(rows=len(table_data), cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.6), Inches(1.2), Inches(2.2), Inches(1.0)]

    for row_idx, row in enumerate(table.rows):
        is_header = (row_idx == 0)
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            val = table_data[row_idx][col_idx]
            cell.text = val
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            run = p.runs[0]
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10)
            
            if is_header:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                set_cell_background(cell, "1F3964")
            else:
                if col_idx == 0:
                    run.bold = True
                if col_idx == 1:
                    run.bold = True
                if row_idx % 2 == 1:
                    set_cell_background(cell, "F2F5F9")
                else:
                    set_cell_background(cell, "FFFFFF")

    # Save to both target locations
    out_path_1 = "AutoRestTest_QFD.docx"
    out_path_2 = os.path.join("resources", "SRS", "AutoRestTest_QFD.docx")
    
    doc.save(out_path_1)
    doc.save(out_path_2)
    print(f"Generated: {out_path_1}")
    print(f"Generated: {out_path_2}")

if __name__ == "__main__":
    create_document()
