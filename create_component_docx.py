import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.color.rgb = RGBColor(0x1F, 0x39, 0x64)  # Navy blue
        if level == 1:
            run.font.size = Pt(16)
            run.bold = True
        elif level == 2:
            run.font.size = Pt(14)
            run.bold = True
        elif level == 3:
            run.font.size = Pt(13)
            run.bold = True
        elif level == 4:
            run.font.size = Pt(12)
            run.bold = True
    heading.paragraph_format.space_after = Pt(6)
    heading.paragraph_format.space_before = Pt(12)

def add_paragraph(doc, text, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.bold = bold
    run.italic = italic
    return p

def set_table_header(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        shading_elm = parse_xml(r'<w:shd {} w:fill="F2F2F2"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading_elm)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10.5)
                r.bold = True

def add_table_row(row, texts):
    for i, text in enumerate(texts):
        cell = row.cells[i]
        cell.text = text
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)

def save_document(doc, path, fallback_path):
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        doc.save(path)
        print(f"Saved: {path}")
        return path
    except PermissionError:
        print(f"Warning: {path} is currently open in Word. Saving to fallback: {fallback_path}")
        doc.save(fallback_path)
        print(f"Saved fallback: {fallback_path}")
        return fallback_path

def main():
    doc = Document()

    # Set page margins to standard 1.0" top/bottom and 1.25" left/right
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    add_heading(doc, '5. Component-Level Design', 1)
    
    # -------------------------------------------------------------------------
    # 5.1 Design Classes and Problem Domain Mapping
    # -------------------------------------------------------------------------
    add_heading(doc, '5.1 Design Classes and Problem Domain Mapping', 2)
    add_paragraph(doc, "AutoRestTest is organized into modular backend services, database models, and intelligent engine components, each addressing a distinct problem domain in automated REST API testing. The table below shows how problem domain concepts map to their corresponding design classes in the codebase and their respective system responsibilities:")
    
    add_paragraph(doc, "Table 5.1: Problem Domain to Design Class Mapping", italic=True)

    table1 = doc.add_table(rows=1, cols=3)
    table1.style = 'Table Grid'
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    table1.autofit = False

    col_widths1 = [Inches(1.5), Inches(2.0), Inches(2.5)]
    for i in range(3):
        table1.columns[i].width = col_widths1[i]
        for cell in table1.columns[i].cells:
            cell.width = col_widths1[i]

    set_table_header(table1.rows[0], ['Domain Concept', 'Design Class', 'Responsibilities'])

    mapping_data = [
        # 1. User & Identity
        ("User Accounts", "User model + AuthService", "User registration, password hashing (Argon2), credential validation, and JWT session token issuance."),
        ("Signup OTP Verification", "PendingSignup model + SignupService", "Pre-registration identity staging, 6-digit email confirmation codes, rate-limiting, and anti-squatting address protection."),
        ("Password Recovery & Tokens", "AuthToken model + TokensService", "Single-use, time-limited security token issuance, validation, and invalidation for password reset workflows."),
        ("Transactional Email Dispatch", "EmailService", "Asynchronous dispatch of verification codes, password reset links, and collaborator invitation notifications."),
        
        # 2. Workspace & Collaboration
        ("Project Workspaces", "Project model + ProjectsService", "Multi-tenant workspace partitioning, project ownership, metadata management, and cross-project data isolation."),
        ("Role-Based Access Control", "ProjectMember model + MembersService + ProjectAccessService", "Granular permission enforcement (admin, tester, viewer) guarding endpoints, test execution, and project settings."),
        ("Collaborator Invitations", "ProjectInvitation model + InvitationsService", "Secure tokenized email invitations, role pre-assignment, and acceptance/revocation lifecycle tracking."),
        
        # 3. Specification & Code Analysis
        ("Specification Ingestion", "ApiSpecification model + SpecsService + SpecificationParser", "OpenAPI 3.0 parsing (JSON/YAML), circular reference resolution (prance), schema validation, and specification updates."),
        ("AI Spec Generation", "SpecGeneration model + SpecGenerationService + GenerationManager", "Source code archive staging, OOPS LLM extraction orchestration, OpenAPI schema generation, and staging review/apply workflow."),
        ("API Endpoints & Schemas", "Endpoint model + EndpointsService", "Route extraction, parameter metadata indexing, manual endpoint injection, and route search/filtering."),
        
        # 4. Dependency Modeling & Caching
        ("Semantic Dependency Graph", "DependencyGraph model + GraphService + OperationGraph", "Directed graph generation mapping API operations to potential data-flow dependencies; supports Focus View topology."),
        ("Semantic Similarity Embedding", "OperationDependencyComparator + EmbeddingModel", "Cosine similarity calculation over property names and schemas using word embeddings to identify producer-consumer relations."),
        ("Graph & State Caching", "CacheConfig + disk cache (shelve)", "Content-addressed caching (spec_<sha256>) of embeddings, graphs, and Q-tables to bypass redundant LLM computation."),
        
        # 5. AI Multi-Agent Testing Engine (MARL)
        ("Test Run Orchestration", "TestSuite model + TestSuitesService + JobManager", "Test execution session management (time budgets, base URLs, route exclusions), engine process spawning, and status polling."),
        ("Multi-Agent RL (MARL)", "AutoRestTest + QLearning (autoresttest-core)", "Value-decomposition Q-learning loop, temporal-difference updates, agent coordination, and epsilon-greedy exploration decay."),
        ("LLM Contextual Value Generation", "SmartValueGenerator + LanguageModel", "Pre-generating realistic, context-aware parameter and body payload values via few-shot structured JSON prompts."),
        ("Dynamic Dependency Data Reuse", "DataSourceAgent + DependencyAgent", "Runtime extraction of identifiers and values from successful 2xx responses and feeding them into downstream dependent requests."),
        ("Request Mutation & Fuzzing", "RequestGenerator (autoresttest-core)", "Assembling HTTP payloads and applying 20% negative mutations (type tampering, missing keys, boundary values) to probe 5xx errors."),
        
        # 6. Execution, Monitoring & Replay
        ("Captured HTTP Request Logging", "TestCase model + RequestLog model", "High-fidelity recording of every sent HTTP request (URL, method, headers, payload, response code, latency, response body)."),
        ("Execution Replay", "TestSuitesService (Replay Mode)", "Deterministic sequential re-execution of previously captured HTTP request sequences without re-calling the AI engine."),
        ("AI Request Descriptions", "RequestDescriptionsService", "Automated batch generation of plain-English summaries explaining the purpose and payload of each executed HTTP test case."),
        
        # 7. Reporting & Administration
        ("Test Analytics & Export", "ReportsService", "Aggregation of operation coverage, status-code distributions, error clustering, and report compilation into PDF and CSV formats."),
        ("AI Failure Root-Cause Analysis", "LlmService (reports module)", "Context-aware LLM evaluation of failed HTTP transactions (request/response context) to explain failure causes in plain language."),
        ("Multi-Scope LLM Configuration", "LlmSettings model + LlmSettingsService", "Dynamic administrative management and runtime hot-reloading of LLM provider keys and models across 3 independent scopes.")
    ]

    for domain, cls, resp in mapping_data:
        row = table1.add_row()
        add_table_row(row, [domain, cls, resp])

    add_paragraph(doc, "")

    # -------------------------------------------------------------------------
    # 5.2 Persistent Data Sources
    # -------------------------------------------------------------------------
    add_heading(doc, '5.2 Persistent Data Sources', 2)
    add_paragraph(doc, "AutoRestTest utilizes distinct persistent data sources to store structured relational data, code archives, embeddings caches, and raw execution logs. The table below outlines these data sources:")

    add_paragraph(doc, "Table 5.2: Persistent Data Sources", italic=True)

    # Clean 3-column table matching Sample Final Report Table 5.2 exactly
    table2 = doc.add_table(rows=1, cols=3)
    table2.style = 'Table Grid'
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    table2.autofit = False

    col_widths2 = [Inches(1.8), Inches(1.4), Inches(2.8)]
    for i in range(3):
        table2.columns[i].width = col_widths2[i]
        for cell in table2.columns[i].cells:
            cell.width = col_widths2[i]

    set_table_header(table2.rows[0], ['Data Source', 'Type', 'Description'])

    sources_data = [
        (
            "PostgreSQL",
            "Relational DB",
            "Primary data store that holds users, projects, specifications, endpoints, test suites, and captured request logs (14 models across Prisma ORM)."
        ),
        (
            "Specification & Code Store",
            "File System",
            "Staging area in uploads/ for incoming OpenAPI schemas (.yaml, .json) and codebase zip archives staged for AI specification generation."
        ),
        (
            "Semantic Graph & Q-Table Cache",
            "Binary Files (shelve)",
            "cache/graphs/ and cache/q_tables/ – Content-addressed disk cache storing NLP embeddings, graph edges, and Q-tables keyed by spec_<sha256>."
        ),
        (
            "Engine Process Logs",
            "Log Files",
            "jobs/<id>/stdout.log – Live stdout/stderr execution and error logs for active MARL testing runs."
        ),
        (
            "Raw Engine Artifacts",
            "JSON Files",
            "data/<spec_name>/ – Detailed engine outputs including report.json, server_errors.json, and operation_status_codes.json."
        )
    ]

    for source, stype, desc in sources_data:
        row = table2.add_row()
        add_table_row(row, [source, stype, desc])

    add_paragraph(doc, "")

    # -------------------------------------------------------------------------
    # 5.3 Behavioral Representation of Core Components
    # -------------------------------------------------------------------------
    add_heading(doc, '5.3 Behavioral Representation of Core Components', 2)
    add_paragraph(doc, "The dynamic behavioral logic of the AutoRestTest platform is centralized around two primary workflows: the generation of OpenAPI specifications from source code, and the end-to-end execution lifecycle of an AI-driven test suite. The following diagrams illustrate the step-by-step state transitions and processes within these core components.")
    
    add_heading(doc, '5.3.1 AI-Driven Testing Workflow', 3)
    add_paragraph(doc, "The automated testing engine operates through a continuous Multi-Agent Reinforcement Learning (MARL) loop. The following activity diagram outlines the execution lifecycle, from the initialization of a new test suite to the compilation of the final report.")
    
    add_paragraph(doc, "Figure 5.2: Complete Testing Workflow Activity Diagram", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    img_path1 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resources", "SRS", "testing_workflow.png")
    if os.path.exists(img_path1):
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_img1 = p_img1.add_run()
        r_img1.add_picture(img_path1, width=Inches(6.0))
    else:
        add_paragraph(doc, "[Image: testing_workflow.png goes here]", italic=True)

    add_heading(doc, '5.3.2 OpenAPI Specification Generation Workflow', 3)
    add_paragraph(doc, "A critical feature of AutoRestTest is the ability to autonomously generate standardized OpenAPI 3.0 specifications from uploaded project source code. The workflow below demonstrates how the OOPS LLM Engine extracts route, controller, and schema information to build the specification.")
    
    add_paragraph(doc, "Figure 5.3: OpenAPI Specification Generation Workflow", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    
    img_path2 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "resources", "SRS", "oas_generation.png")
    if os.path.exists(img_path2):
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_img2 = p_img2.add_run()
        r_img2.add_picture(img_path2, width=Inches(6.0))
    else:
        add_paragraph(doc, "[Image: oas_generation.png goes here]", italic=True)

    add_paragraph(doc, "")

    # Save to root and resources/SRS
    output_path_root = os.path.join("D:\\SPL3\\AutoRestTest", "AutoRestTest_ComponentLevelDesign.docx")
    fallback_root = os.path.join("D:\\SPL3\\AutoRestTest", "AutoRestTest_ComponentLevelDesign_Updated.docx")
    
    output_path_srs = os.path.join("D:\\SPL3\\AutoRestTest", "resources", "SRS", "AutoRestTest_ComponentLevelDesign.docx")
    fallback_srs = os.path.join("D:\\SPL3\\AutoRestTest", "resources", "SRS", "AutoRestTest_ComponentLevelDesign_Updated.docx")

    save_document(doc, output_path_root, fallback_root)
    save_document(doc, output_path_srs, fallback_srs)

if __name__ == "__main__":
    main()
